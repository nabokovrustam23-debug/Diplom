#!/usr/bin/env python3
# coding: utf-8
"""
Сборка пояснительной записки ВКР в формате .docx по шаблонам ККЭП.

Использует:
  - Официальные шаблоны ККЭП из docs/templates_kkep/Шаблоны_ВКР_09.00.00_2026_/
  - Black-box markdown-главы из docs/pz/*.md
  - pandoc для рендера тела
  - python-docx + docxcompose для сшивки

Запуск:
  python3 docs/build_vkr.py

Результат:
  /home/ubuntu/repos/Diplom/DIPLOM_ВКР_Тихий_час.docx
"""

from __future__ import annotations

import os
import shutil
import subprocess
import sys
from copy import deepcopy
from pathlib import Path

import docx
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt
from docxcompose.composer import Composer

ROOT = Path(__file__).resolve().parent.parent
TPL_DIR = ROOT / "docs" / "templates_kkep" / "Шаблоны_ВКР_09.00.00_2026_"
PZ_DIR = ROOT / "docs" / "pz"
BUILD_DIR = ROOT / "build"

THEME = (
    "Разработка модулей системы управления взаимоотношениями "
    "с клиентами сети барбершопов «Тихий час»"
)
GROUP = "67-Д9-4ИСП"
YEAR = "2026"
PLACEHOLDER = "_____________"
SHIFR = "ККЭП 09.02.07 ____ ПЗ"

PZ_FILES = [
    "00-vvedenie.md",
    "01-naznachenie.md",
    "02-analiz.md",
    "03-realizatsiya.md",
    "04-testirovanie.md",
    "05-ekspluatatsiya.md",
    "06-zaklyuchenie.md",
    "07-istochniki.md",
    "prilozhenie-a.md",
    "prilozhenie-b.md",
    "prilozhenie-v.md",
]


# ---------------------------------------------------------------------------
# helpers


def _norm(s):
    return s.replace("\u00a0", " ").replace("\u2009", " ")


def replace_text_in_runs(paragraph, old, new):
    """Заменить текст в параграфе с учётом разбиения на runs.
    Нормализует неразрывные пробелы при сравнении."""
    full = "".join(r.text for r in paragraph.runs)
    if _norm(old) not in _norm(full):
        return False
    # Найдём позицию вхождения в нормализованной строке и сделаем замену
    # на нормализованной полной строке.
    norm_full = _norm(full)
    idx = norm_full.find(_norm(old))
    # Простейшая стратегия: построить новую полную строку и положить в run 0.
    # Сохраняет шрифт первого run-а — этого достаточно для штампов.
    # Используем нормализованную версию full, чтобы избежать проблем с nbsp.
    new_full = norm_full[:idx] + new + norm_full[idx + len(_norm(old)):]
    if not paragraph.runs:
        paragraph.add_run(new_full)
        return True
    paragraph.runs[0].text = new_full
    for r in paragraph.runs[1:]:
        r.text = ""
    return True


def replace_in_doc(doc, mapping):
    """Заменить плейсхолдеры в теле, header/footer, таблицах документа."""

    def walk_paragraphs(container):
        for p in container.paragraphs:
            for old, new in mapping.items():
                replace_text_in_runs(p, old, new)
        for tbl in container.tables:
            for row in tbl.rows:
                for cell in row.cells:
                    walk_paragraphs(cell)

    walk_paragraphs(doc)
    for s in doc.sections:
        walk_paragraphs(s.first_page_header)
        walk_paragraphs(s.header)
        walk_paragraphs(s.first_page_footer)
        walk_paragraphs(s.footer)


def clear_body(doc):
    """Удалить весь body, оставив sectPr."""
    body = doc.element.body
    keep_sect_pr = None
    to_remove = []
    for child in body:
        if child.tag == qn("w:sectPr"):
            keep_sect_pr = child
            continue
        to_remove.append(child)
    for c in to_remove:
        body.remove(c)


def ensure_heading_styles(doc):
    """Задать стили Heading 1/2/3 в стиле ГОСТ для ВКР."""
    styles = doc.styles
    spec = [
        ("Heading 1", True, WD_ALIGN_PARAGRAPH.LEFT, True),
        ("Heading 2", True, WD_ALIGN_PARAGRAPH.LEFT, False),
        ("Heading 3", True, WD_ALIGN_PARAGRAPH.LEFT, False),
    ]
    for name, bold, align, page_break in spec:
        try:
            st = styles[name]
        except KeyError:
            st = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st.font.name = "Times New Roman"
        # Также прописать восточно-азиатский / cs шрифт
        rPr = st.element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.insert(0, rFonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rFonts.set(qn(attr), "Times New Roman")
        st.font.size = Pt(14)
        st.font.bold = bold
        st.font.color.rgb = None  # авто
        pf = st.paragraph_format
        pf.alignment = align
        pf.first_line_indent = Cm(1.25)
        pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.page_break_before = page_break
        # Убрать наследование verifyText от Normal на лишний интервал
    # Normal
    try:
        normal = styles["Normal"]
        normal.font.name = "Times New Roman"
        normal.font.size = Pt(14)
        rPr = normal.element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.insert(0, rFonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rFonts.set(qn(attr), "Times New Roman")
        normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        normal.paragraph_format.first_line_indent = Cm(1.25)
        normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        normal.paragraph_format.space_before = Pt(0)
        normal.paragraph_format.space_after = Pt(0)
    except KeyError:
        pass


# ---------------------------------------------------------------------------
# steps


def make_reference_docx(out_path: Path):
    """Создать reference.docx из шаблона «Содержание и последующие листы»
    с очищенным body, корректными стилями и плейсхолдерами штампа."""
    src = TPL_DIR / "Содержание и последующие листы_ПРИМЕР_2026_.docx"
    d = docx.Document(str(src))
    clear_body(d)
    ensure_heading_styles(d)
    replace_in_doc(d, get_stamp_mapping())
    # Нужен хотя бы один параграф, чтобы pandoc подхватил
    p = d.add_paragraph("")
    d.save(str(out_path))
    _fix_doc_defaults(out_path)


def _fix_doc_defaults(docx_path: Path):
    """Перезаписать docDefaults в styles.xml на Times New Roman 14pt + 1.5 line."""
    import zipfile, shutil, tempfile, re

    new_defaults = (
        '<w:docDefaults>'
        '<w:rPrDefault><w:rPr>'
        '<w:rFonts w:ascii="Times New Roman" w:eastAsia="Times New Roman" '
        'w:hAnsi="Times New Roman" w:cs="Times New Roman"/>'
        '<w:sz w:val="28"/>'  # 14pt = 28 half-points
        '<w:szCs w:val="28"/>'
        '<w:lang w:val="ru-RU" w:eastAsia="en-US" w:bidi="ar-SA"/>'
        '</w:rPr></w:rPrDefault>'
        '<w:pPrDefault><w:pPr>'
        '<w:spacing w:after="0" w:line="360" w:lineRule="auto"/>'  # 1.5 line spacing (240*1.5=360)
        '<w:ind w:firstLine="709"/>'  # 1.25 cm (709 twips)
        '<w:jc w:val="both"/>'
        '</w:pPr></w:pPrDefault>'
        '</w:docDefaults>'
    )

    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp) / "patched.docx"
        with zipfile.ZipFile(docx_path, "r") as zin, zipfile.ZipFile(
            tmp_path, "w", zipfile.ZIP_DEFLATED
        ) as zout:
            for item in zin.namelist():
                data = zin.read(item)
                if item == "word/styles.xml":
                    text = data.decode("utf-8")
                    text = re.sub(
                        r'<w:docDefaults>.*?</w:docDefaults>',
                        new_defaults,
                        text,
                        count=1,
                        flags=re.DOTALL,
                    )
                    data = text.encode("utf-8")
                zout.writestr(item, data)
        shutil.move(str(tmp_path), str(docx_path))


def get_stamp_mapping():
    return {
        "ККЭП 09.02.01 2222 ПЗ": SHIFR,
        "Разработка LED-подсветки для экрана": THEME,
        "Смоленова": PLACEHOLDER,
        "Иванов": PLACEHOLDER,
        "Петров": PLACEHOLDER,
        "Сторчак": PLACEHOLDER,
        "Головко": PLACEHOLDER,
        "Гр. 555-Д9-4КСК": f"Гр. {GROUP}",
    }


def make_titul(out_path: Path):
    src = TPL_DIR / "Титульный лист ИНС  ИСП  Дневное_2026_.docx"
    d = docx.Document(str(src))
    mapping = {
        "Разработка LED-подсветки для экрана": THEME,
        "С.В. Смоленова": PLACEHOLDER,
        "541-Д9-4ИНС": GROUP,
        "И.И. Иванов": PLACEHOLDER,
        "П.П. Петров": PLACEHOLDER,
        "Е.В. Лебедь": PLACEHOLDER,
    }
    replace_in_doc(d, mapping)
    d.save(str(out_path))


def make_zadanie(out_path: Path):
    src = TPL_DIR / "Лист задания без граф. части_2026_.docx"
    d = docx.Document(str(src))
    # Заменим только основные данные: ФИО студента и тему ВКР.
    # Список разделов ПЗ оставим как есть — руководитель скорректирует
    # его в финальной версии под фактическое содержание работы.
    mapping = {
        "Сидоровой Анне Сергеевне": PLACEHOLDER,
        "Разработка веб-системы для клиентов компании": THEME,
        "ООО «ИмпульСС»": "",
    }
    replace_in_doc(d, mapping)
    replace_in_doc(d, get_stamp_mapping())
    d.save(str(out_path))


def get_annotation_body():
    """Текст аннотации с подставленной темой."""
    return [
        "Аннотация",
        "",
        f"Темой выпускной квалификационной работы является {THEME[0].lower()+THEME[1:]}.",
        "",
        "Результаты выполнения выпускной квалификационной работы представлены "
        "в виде пояснительной записки и разработанного программно-информационного "
        "комплекса, реализующего CRM-систему сети барбершопов с поддержкой "
        "многофилиальной структуры и онлайн-записи на услугу.",
        "",
        "В пояснительной записке изложены: анализ предметной области сети "
        "барбершопов и постановка задачи; обоснование выбора средств разработки "
        "(ASP.NET Core 8, MS SQL Server, Razor Pages, Bootstrap-альтернатива в виде "
        "собственной дизайн-системы); проектирование модели данных по паттерну "
        "Persona; реализация модулей онлайн-записи на услугу, расчёта свободных "
        "слотов, административной панели владельца сети; модульное тестирование "
        "сервисного слоя; руководство пользователя. В составе приложений приведены "
        "техническое задание, основные экранные формы и программный код, "
        "тест-кейсы.",
        "",
        "Пояснительная записка состоит из ___ листов, ___ изображений, "
        "___ таблиц и 3 приложений, которые в совокупности отображают результаты "
        "всех этапов разработки CRM-системы. При разработке использовались "
        "методы объектно-ориентированного проектирования, паттерны слоистой "
        "архитектуры (Layered Architecture), принципы SOLID, REST-подобный подход "
        "к организации внутренних веб-сервисов на Razor Pages.",
    ]


def make_opis_annotatsiya(out_path: Path):
    src = TPL_DIR / "Опись  Аннотация ИСП гр67_75_Дневное без граф. части_2026_.docx"
    d = docx.Document(str(src))
    # Заменить шифр в описи и штампе.
    replace_in_doc(d, {"ККЭП 09.02.07 0142 ПЗ": SHIFR})
    replace_in_doc(d, get_stamp_mapping())

    # Заменить тестовый текст аннотации на наш.
    # Найдём первый параграф со словом «Темой» и далее — всё перепишем.
    body = d.element.body
    paras = list(d.paragraphs)
    start_i = None
    for i, p in enumerate(paras):
        if p.text.strip().startswith("Темой выпускной"):
            start_i = i
            break
    if start_i is not None:
        # Удалим старые «текстовые» параграфы (не таблицы и не пустые перед началом).
        end_i = len(paras)
        for p in paras[start_i:end_i]:
            elem = p._element
            elem.getparent().remove(elem)
        # Вставим наши параграфы аннотации.
        new_paras_text = get_annotation_body()
        # Первый «Аннотация» уже есть как заголовок выше; вставим только тело.
        # Пропустим первые 2 элемента (заголовок и пустую) если они уже есть.
        for line in new_paras_text[2:]:
            np = d.add_paragraph(line)
            np.paragraph_format.first_line_indent = Cm(1.25)
            np.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            np.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            for r in np.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(14)

    d.save(str(out_path))


def make_normokontrol(out_path: Path):
    src = TPL_DIR / "__Лист нормоконтроля ВКР_2026_.docx"
    d = docx.Document(str(src))
    # Заменим тестовый текст шаблона на нашу тему/группу.
    # В шаблоне тема описана 16 словами «тема» с двумя заглавными.
    fake_theme = "Тема тема тема тема тема Тема тема тема тема тема тема тема тема тема тема тема"
    fake_theme_tail = "тема тема тема тема тема тема тема тема тема тема"
    mapping = {
        fake_theme: THEME,
        fake_theme_tail: "",
        "тема тема тема тема тема тема тема тема": "",
        "тема тема тема тема тема": "",
        "тема тема тема": "",
        "тема тема": "",
        "54 – КД9 – 4 ИСП": GROUP,
        "09.02.07 Информационные системы и программирование, программист":
            "09.02.07 Информационные системы и программирование",
        "Петрова Петра Петровича": PLACEHOLDER,
        "_____________а Петра _____________ича": PLACEHOLDER,
    }
    replace_in_doc(d, mapping)
    replace_in_doc(d, get_stamp_mapping())
    d.save(str(out_path))


# ---------------------------------------------------------------------------
# pandoc body


def merge_markdown(out_path: Path):
    parts = []
    for fname in PZ_FILES:
        f = PZ_DIR / fname
        if not f.exists():
            print(f"  WARN: {f} not found, skipping", file=sys.stderr)
            continue
        text = f.read_text(encoding="utf-8")
        if parts:
            parts.append("\n\n\\newpage\n\n")
        parts.append(text)
    out_path.write_text("\n".join(parts), encoding="utf-8")


def render_body(md_path: Path, ref_path: Path, out_path: Path):
    cmd = [
        "pandoc",
        str(md_path),
        "--reference-doc",
        str(ref_path),
        "--resource-path",
        str(PZ_DIR),
        "--toc",
        "--toc-depth=2",
        "-f",
        "markdown",
        "-t",
        "docx",
        "-o",
        str(out_path),
    ]
    print(" ".join(cmd))
    subprocess.run(cmd, check=True)
    # Pandoc 2.9 хардкодит "Table of Contents" в заголовок TOC и переводит его
    # на язык документа только в LaTeX/HTML. Заменим вручную в document.xml.
    _replace_toc_title(out_path, "Table of Contents", "Содержание")


def _replace_toc_title(docx_path: Path, old: str, new: str):
    """Заменить текст в word/document.xml внутри готового docx."""
    import zipfile, shutil, tempfile
    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp) / "patched.docx"
        with zipfile.ZipFile(docx_path, "r") as zin, zipfile.ZipFile(
            tmp_path, "w", zipfile.ZIP_DEFLATED
        ) as zout:
            for item in zin.namelist():
                data = zin.read(item)
                if item == "word/document.xml":
                    text = data.decode("utf-8")
                    text = text.replace(old, new)
                    data = text.encode("utf-8")
                zout.writestr(item, data)
        shutil.move(str(tmp_path), str(docx_path))


def add_soderzhanie_marker(md_path: Path):
    """Заголовок 'Содержание' добавим pandoc через --toc, тут пока ничего не делаем."""
    return


# ---------------------------------------------------------------------------
# compose


def compose(parts, out_path: Path):
    master = docx.Document(str(parts[0]))
    composer = Composer(master)
    for p in parts[1:]:
        composer.append(docx.Document(str(p)))
    composer.save(str(out_path))


# ---------------------------------------------------------------------------


def main():
    BUILD_DIR.mkdir(exist_ok=True, parents=True)

    titul = BUILD_DIR / "01_titul.docx"
    zadanie = BUILD_DIR / "02_zadanie.docx"
    opis = BUILD_DIR / "03_opis_annotatsiya.docx"
    reference = BUILD_DIR / "_reference.docx"
    merged_md = BUILD_DIR / "_merged.md"
    body = BUILD_DIR / "04_body.docx"
    normokontrol = BUILD_DIR / "05_normokontrol.docx"

    final_path = ROOT / "DIPLOM_ВКР_Тихий_час.docx"

    print("[1/7] Титульный лист...")
    make_titul(titul)
    print("[2/7] Лист задания...")
    make_zadanie(zadanie)
    print("[3/7] Опись + аннотация...")
    make_opis_annotatsiya(opis)
    print("[4/7] reference.docx (шаблон с рамкой)...")
    make_reference_docx(reference)
    print("[5/7] merged.md...")
    merge_markdown(merged_md)
    add_soderzhanie_marker(merged_md)
    print("[6/7] pandoc body...")
    render_body(merged_md, reference, body)
    print("[7/7] Лист нормоконтроля...")
    make_normokontrol(normokontrol)

    print("Сшивка...")
    compose([titul, zadanie, opis, body, normokontrol], final_path)
    print(f"Готово: {final_path} ({final_path.stat().st_size / 1024:.1f} KB)")


if __name__ == "__main__":
    main()
